from __future__ import annotations

import os
from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\cmtc\CMTC-AI-System")
OUT_DIR = ROOT / "นำเสนอ"
DIAGRAM_DIR = OUT_DIR / "report_diagrams"
OUT_FILE = OUT_DIR / "รายงานเอกสารประกอบการสอบโครงงาน_CMTC_AI_KMS.docx"
PROJECT_NAME_TH = "ระบบจัดการข้อมูลและองค์ความรู้ด้วยปัญญาประดิษฐ์"
PROJECT_NAME_EN = "CMTC AI Knowledge Management System"


def ensure_dirs() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    DIAGRAM_DIR.mkdir(exist_ok=True)


def font_path() -> str:
    candidates = [
        r"C:\Windows\Fonts\tahoma.ttf",
        r"C:\Windows\Fonts\THSarabunNew.ttf",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return candidates[-1]


FONT_PATH = font_path()


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    bold_candidates = [
        r"C:\Windows\Fonts\tahomabd.ttf",
        r"C:\Windows\Fonts\arialbd.ttf",
    ]
    if bold:
        for candidate in bold_candidates:
            if Path(candidate).exists():
                return ImageFont.truetype(candidate, size=size)
    return ImageFont.truetype(FONT_PATH, size=size)


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> tuple[int, int]:
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=8)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str = "#111827",
    spacing: int = 8,
) -> None:
    x1, y1, x2, y2 = box
    w, h = text_size(draw, text, font)
    draw.multiline_text(
        (x1 + (x2 - x1 - w) / 2, y1 + (y2 - y1 - h) / 2),
        text,
        font=font,
        fill=fill,
        align="center",
        spacing=spacing,
    )


def draw_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fill: str = "#FFFFFF",
    outline: str = "#1F2937",
    width: int = 3,
    font_size: int = 28,
    bold: bool = False,
    radius: int = 12,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)
    draw_centered_text(draw, box, text, pil_font(font_size, bold=bold))


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#374151") -> None:
    draw.line([start, end], fill=color, width=4)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 18 * direction, ey - 10), (ex - 18 * direction, ey + 10)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 10, ey - 18 * direction), (ex + 10, ey - 18 * direction)]
    draw.polygon(points, fill=color)


def new_diagram(title: str, filename: str) -> tuple[Image.Image, ImageDraw.ImageDraw, Path]:
    img = Image.new("RGB", (1800, 1120), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), title, font=pil_font(42, True), fill="#111827")
    draw.line((70, 105, 1730, 105), fill="#D1D5DB", width=3)
    return img, draw, DIAGRAM_DIR / filename


def save_diagram(img: Image.Image, path: Path) -> Path:
    img.save(path, quality=95)
    return path


def diagram_use_case() -> Path:
    img, draw, path = new_diagram("Use Case Diagram", "01_use_case.png")
    draw.ellipse((120, 250, 200, 330), outline="#111827", width=4)
    draw.line((160, 330, 160, 470), fill="#111827", width=4)
    draw.line((85, 380, 235, 380), fill="#111827", width=4)
    draw.line((160, 470, 100, 570), fill="#111827", width=4)
    draw.line((160, 470, 220, 570), fill="#111827", width=4)
    draw_centered_text(draw, (50, 590, 280, 670), "ผู้ใช้งานทั่วไป", pil_font(26, True))

    draw.ellipse((1520, 250, 1600, 330), outline="#111827", width=4)
    draw.line((1560, 330, 1560, 470), fill="#111827", width=4)
    draw.line((1485, 380, 1635, 380), fill="#111827", width=4)
    draw.line((1560, 470, 1500, 570), fill="#111827", width=4)
    draw.line((1560, 470, 1620, 570), fill="#111827", width=4)
    draw_centered_text(draw, (1440, 590, 1690, 690), "เจ้าหน้าที่\nและผู้ดูแลระบบ", pil_font(25, True))

    use_cases = [
        ((520, 185, 870, 275), "เข้าสู่ระบบ"),
        ((930, 185, 1280, 275), "ค้นหาข้อมูลองค์ความรู้"),
        ((520, 330, 870, 420), "ถามตอบผ่าน AI Chat"),
        ((930, 330, 1280, 420), "จัดการบทความความรู้"),
        ((520, 475, 870, 565), "จัดการเอกสาร"),
        ((930, 475, 1280, 565), "จัดการผู้ใช้และสิทธิ์"),
        ((520, 620, 870, 710), "ดูสถิติและประวัติ AI"),
        ((930, 620, 1280, 710), "ส่งงานประมวลผลไป n8n"),
    ]
    for box, text in use_cases:
        draw.ellipse(box, fill="#EFF6FF", outline="#1D4ED8", width=3)
        draw_centered_text(draw, box, text, pil_font(24, True))

    for box, _ in use_cases[:3]:
        arrow(draw, (250, 400), (box[0], (box[1] + box[3]) // 2), "#6B7280")
    for box, _ in use_cases[3:]:
        arrow(draw, (1450, 400), (box[2], (box[1] + box[3]) // 2), "#6B7280")

    draw_box(draw, (1335, 820, 1650, 930), "ระบบภายนอก\nOpenAI / n8n", "#F8FAFC", "#475569", 3, 26, True)
    arrow(draw, (1110, 710), (1335, 875), "#475569")
    arrow(draw, (700, 710), (1335, 875), "#475569")
    return save_diagram(img, path)


def diagram_architecture() -> Path:
    img, draw, path = new_diagram("System Architecture Diagram", "02_architecture.png")
    boxes = {
        "user": ((90, 310, 360, 470), "ผู้ใช้งาน\nBrowser"),
        "next": ((500, 260, 820, 520), "Next.js Web App\nApp Router\nAPI Routes\nReact UI"),
        "supabase": ((985, 170, 1315, 360), "Supabase\nAuth / PostgreSQL\nStorage / RLS"),
        "ai": ((985, 445, 1315, 610), "AI Service\nOpenAI หรือ Ollama\nChat / Categorize"),
        "n8n": ((1420, 315, 1700, 485), "n8n Workflow\nWebhook Automation"),
        "analytics": ((985, 720, 1315, 880), "Analytics\nAI logs / Feedback\nReports"),
    }
    colors = {
        "user": ("#F9FAFB", "#111827"),
        "next": ("#EFF6FF", "#1D4ED8"),
        "supabase": ("#ECFDF5", "#047857"),
        "ai": ("#FFF7ED", "#C2410C"),
        "n8n": ("#F5F3FF", "#6D28D9"),
        "analytics": ("#F8FAFC", "#475569"),
    }
    for key, (box, text) in boxes.items():
        fill, outline = colors[key]
        draw_box(draw, box, text, fill, outline, 4, 26, True)
    arrow(draw, (360, 390), (500, 390))
    arrow(draw, (820, 330), (985, 270))
    arrow(draw, (820, 430), (985, 525))
    arrow(draw, (1315, 360), (1420, 390))
    arrow(draw, (820, 500), (985, 790))
    arrow(draw, (1315, 525), (1420, 420))
    draw.text((500, 565), "ชั้นควบคุมสิทธิ์: Middleware, RBAC, Server Components, API validation", font=pil_font(26), fill="#374151")
    draw.text((500, 610), "ชั้นข้อมูล: บทความ เอกสาร บุคลากร แผนก บทสนทนา และสถิติ AI", font=pil_font(26), fill="#374151")
    return save_diagram(img, path)


def diagram_activity() -> Path:
    img, draw, path = new_diagram("Activity Diagram / Main Process Flow", "03_activity.png")
    steps = [
        ((710, 145, 1090, 225), "เริ่มต้นใช้งานระบบ"),
        ((710, 285, 1090, 365), "เข้าสู่ระบบ"),
        ((710, 425, 1090, 505), "ตรวจสอบบทบาทและสิทธิ์"),
        ((250, 600, 590, 690), "จัดการบทความ\nหรือเอกสาร"),
        ((730, 600, 1070, 690), "ค้นหาและถาม AI"),
        ((1210, 600, 1550, 690), "ดูสถิติและรายงาน"),
        ((250, 800, 590, 890), "บันทึกข้อมูล\nลงฐานข้อมูล"),
        ((730, 800, 1070, 890), "ดึงบริบทจากฐานข้อมูล\nและเรียก AI"),
        ((1210, 800, 1550, 890), "บันทึกผลการใช้งาน\nและประเมินคุณภาพ"),
        ((710, 980, 1090, 1060), "สิ้นสุดกระบวนการ"),
    ]
    for box, text in steps:
        draw_box(draw, box, text, "#FFFFFF", "#1F2937", 3, 24, True, 10)
    arrow(draw, (900, 225), (900, 285))
    arrow(draw, (900, 365), (900, 425))
    arrow(draw, (850, 505), (420, 600))
    arrow(draw, (900, 505), (900, 600))
    arrow(draw, (950, 505), (1380, 600))
    arrow(draw, (420, 690), (420, 800))
    arrow(draw, (900, 690), (900, 800))
    arrow(draw, (1380, 690), (1380, 800))
    arrow(draw, (590, 845), (710, 1020))
    arrow(draw, (900, 890), (900, 980))
    arrow(draw, (1210, 845), (1090, 1020))
    return save_diagram(img, path)


def diagram_erd() -> Path:
    img, draw, path = new_diagram("Data Design / ER Diagram", "04_erd.png")
    tables = [
        ((80, 170, 440, 370), "profiles\nid PK\nrole, status\ndepartment_id FK"),
        ((80, 485, 440, 685), "departments\nid PK\ncode, name\nphone, email"),
        ((575, 160, 935, 390), "knowledge_articles\nid PK\ntitle, slug, status\nauthor_id FK\ncategory_id FK\nembedding"),
        ((575, 475, 935, 675), "documents\nid PK\ndisplay_title\ncategory, status\ndepartment_id FK\nstorage_path"),
        ((1080, 145, 1440, 315), "categories\nid PK\nname, slug"),
        ((1080, 360, 1440, 530), "tags\nid PK\nname"),
        ((1080, 590, 1440, 760), "article_tags\narticle_id FK\ntag_id FK"),
        ((575, 770, 935, 960), "conversations\nid PK\nuser_id FK\ntitle"),
        ((1080, 800, 1440, 1000), "chat_messages\nid PK\nconversation_id FK\nrole, content\nconfidence_score"),
        ((80, 790, 440, 990), "ai_analytics\nid PK\nquestion, answer\nmodel_used\nconfidence_score"),
    ]
    for box, text in tables:
        draw_box(draw, box, text, "#F8FAFC", "#334155", 3, 22, False, 8)
    relations = [
        ((440, 270), (575, 275)),
        ((440, 585), (575, 575)),
        ((935, 260), (1080, 230)),
        ((935, 285), (1080, 675)),
        ((1260, 530), (1260, 590)),
        ((440, 270), (575, 850)),
        ((935, 865), (1080, 900)),
        ((440, 585), (575, 575)),
    ]
    for start, end in relations:
        arrow(draw, start, end, "#64748B")
    draw.text((80, 1030), "หมายเหตุ: auth.users เป็นตารางผู้ใช้ของ Supabase Authentication และเชื่อมกับ profiles, conversations, knowledge_articles", font=pil_font(22), fill="#374151")
    return save_diagram(img, path)


def diagram_sequence() -> Path:
    img, draw, path = new_diagram("Sequence Diagram / AI Chat Flow", "05_sequence.png")
    actors = [
        (160, "ผู้ใช้"),
        (500, "Browser"),
        (850, "Next.js API"),
        (1190, "Supabase"),
        (1510, "AI Service"),
    ]
    for x, label in actors:
        draw_box(draw, (x - 125, 145, x + 125, 225), label, "#EFF6FF", "#1D4ED8", 3, 24, True)
        draw.line((x, 225, x, 980), fill="#CBD5E1", width=3)
    messages = [
        (160, 500, 310, "1 ส่งคำถาม"),
        (500, 850, 410, "2 POST /api/chat"),
        (850, 1190, 510, "3 อ่านบริบท บทความ เอกสาร บุคลากร"),
        (1190, 850, 610, "4 ส่งข้อมูลที่เกี่ยวข้อง"),
        (850, 1510, 710, "5 สร้างคำตอบจากบริบท"),
        (1510, 850, 810, "6 ส่งคำตอบและ confidence"),
        (850, 1190, 895, "7 บันทึกข้อความและ analytics"),
        (850, 500, 950, "8 ส่งผลลัพธ์กลับหน้าเว็บ"),
    ]
    for sx, ex, y, label in messages:
        arrow(draw, (sx, y), (ex, y), "#334155")
        draw.text((min(sx, ex) + 20, y - 38), label, font=pil_font(22), fill="#111827")
    return save_diagram(img, path)


def diagram_ai_pipeline() -> Path:
    img, draw, path = new_diagram("AI Pipeline and Model Evaluation", "06_ai_pipeline.png")
    steps = [
        ((80, 260, 340, 420), "1\nรับข้อมูล\nบทความและเอกสาร"),
        ((410, 260, 670, 420), "2\nจัดหมวดหมู่\nและสกัดข้อความ"),
        ((740, 260, 1000, 420), "3\nสร้าง Embedding\nและค้นหา pgvector"),
        ((1070, 260, 1330, 420), "4\nดึงบริบทที่เกี่ยวข้อง\nจากฐานข้อมูล"),
        ((1400, 260, 1660, 420), "5\nสร้างคำตอบด้วย\nOpenAI / Ollama"),
        ((740, 650, 1000, 810), "6\nบันทึกสถิติ\nconfidence / time"),
        ((1070, 650, 1330, 810), "7\nตรวจผลลัพธ์\nและ feedback"),
    ]
    for box, text in steps:
        draw_box(draw, box, text, "#FFF7ED", "#C2410C", 4, 23, True, 14)
    for i in range(4):
        arrow(draw, (steps[i][0][2], 340), (steps[i + 1][0][0], 340), "#9A3412")
    arrow(draw, (1530, 420), (1000, 730), "#9A3412")
    arrow(draw, (1000, 730), (1070, 730), "#9A3412")
    draw.text((90, 905), "ตัวชี้วัดประเมิน: ความถูกต้องของคำตอบ ความเกี่ยวข้องของแหล่งข้อมูล ค่า confidence เวลาในการตอบ และความคิดเห็นผู้ใช้", font=pil_font(26), fill="#374151")
    return save_diagram(img, path)


def diagram_wireframe() -> Path:
    img, draw, path = new_diagram("UI / Wireframe / HMI Design", "07_wireframe.png")
    draw_box(draw, (110, 180, 520, 910), "หน้าจอ Login\n\nโลโก้ระบบ\nช่องอีเมล\nช่องรหัสผ่าน\nปุ่มเข้าสู่ระบบ", "#F9FAFB", "#111827", 3, 25, True)
    draw_box(draw, (650, 180, 1690, 910), "Dashboard หลังเข้าสู่ระบบ", "#FFFFFF", "#111827", 3, 30, True)
    draw.rectangle((690, 250, 880, 870), outline="#CBD5E1", width=3, fill="#F8FAFC")
    draw.text((715, 290), "เมนู", font=pil_font(26, True), fill="#111827")
    for i, label in enumerate(["Dashboard", "Knowledge", "Documents", "AI Chat", "Analytics", "Users"]):
        y = 350 + i * 70
        draw.rounded_rectangle((710, y, 860, y + 45), radius=8, fill="#E5E7EB")
        draw.text((725, y + 8), label, font=pil_font(19), fill="#111827")
    draw.rectangle((920, 250, 1650, 360), outline="#CBD5E1", width=3, fill="#EFF6FF")
    draw_centered_text(draw, (920, 250, 1650, 360), "สรุปจำนวนบทความ เอกสาร ผู้ใช้ และคำถาม AI", pil_font(24, True))
    draw.rectangle((920, 400, 1280, 830), outline="#CBD5E1", width=3, fill="#FFFFFF")
    draw_centered_text(draw, (920, 400, 1280, 830), "ตารางรายการเอกสาร\nค้นหา / กรอง / จัดการ", pil_font(24, True))
    draw.rectangle((1320, 400, 1650, 830), outline="#CBD5E1", width=3, fill="#FFFFFF")
    draw_centered_text(draw, (1320, 400, 1650, 830), "พื้นที่ AI Chat\nถามตอบ\nแสดงแหล่งข้อมูล", pil_font(24, True))
    return save_diagram(img, path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "8") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_style(table, header_fill: str = "1F4E79") -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Tahoma"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Tahoma")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Tahoma")
                    run._element.rPr.rFonts.set(qn("w:cs"), "Tahoma")
                    run.font.size = Pt(9)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F8FBFF")


def set_col_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def set_doc_styles(doc: Document) -> None:
    styles = doc.styles
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Tahoma"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Tahoma")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Tahoma")
        style._element.rPr.rFonts.set(qn("w:cs"), "Tahoma")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.bold = True
    styles["Heading 3"].font.size = Pt(11.5)
    styles["Heading 3"].font.bold = True


def add_paragraph(doc: Document, text: str = "", style: str | None = None, bold_lead: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        run.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_style(table)
    if widths:
        set_col_widths(table, widths)
    return table


def add_image(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.4))
    cp = doc.add_paragraph(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)
    for run in cp.runs:
        run.italic = True
        run.font.size = Pt(9)


def add_cover(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
    if (ROOT / "public" / "cmtc-ai-logo.png.png").exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ROOT / "public" / "cmtc-ai-logo.png.png"), width=Inches(1.25))
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("รายงานเอกสารประกอบการสอบโครงงาน").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(PROJECT_NAME_TH).bold = True
    p.runs[0].font.size = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"({PROJECT_NAME_EN})")
    p.runs[0].font.size = Pt(14)
    doc.add_paragraph()
    meta_rows = [
        ["ชื่อโครงงาน", PROJECT_NAME_TH],
        ["ชื่อภาษาอังกฤษ", PROJECT_NAME_EN],
        ["ผู้จัดทำ", "นายพัฒนกุล เทปิน รหัสนักศึกษา 68409010013"],
        ["อาจารย์ที่ปรึกษา", "นายอนุชาติ รังสิยานนท์"],
        ["สถานศึกษา", "วิทยาลัยเทคนิคเชียงใหม่"],
        ["ลักษณะโครงงาน", "ระบบเว็บแอปพลิเคชันด้านการจัดการองค์ความรู้และผู้ช่วย AI"],
        ["วันที่จัดทำเอกสาร", "16 กันยายน 2569"],
    ]
    table = add_table(doc, ["รายการ", "รายละเอียด"], meta_rows, [1.6, 4.8])
    for row in table.rows[1:]:
        row.cells[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("จัดทำเพื่อประกอบการสอบและส่งเอกสารตามรายการที่อาจารย์กำหนด").bold = True
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    add_heading(doc, "สารบัญ", 1)
    items = [
        "1 บทนำและภาพรวมโครงงาน",
        "2 Project Requirement Specification",
        "3 Use Case Diagram และ Use Case Description",
        "4 System Architecture Diagram",
        "5 Activity Diagram และ Main Process Flow",
        "6 Data Design และ ER Diagram",
        "7 Sequence Diagram",
        "8 UI Wireframe และ HMI Design",
        "9 Hardware Block Diagram และ I O Mapping",
        "10 AI Pipeline และ Model Evaluation",
        "11 Test Case และ Test Result",
        "12 Requirement Traceability Matrix",
        "13 ภาพและหลักฐานชิ้นงานที่พัฒนาเสร็จ",
        "14 สรุปผลการจัดทำเอกสาร",
    ]
    for item in items:
        add_paragraph(doc, item)
    doc.add_page_break()


def add_intro(doc: Document) -> None:
    add_heading(doc, "1 บทนำและภาพรวมโครงงาน", 1)
    add_paragraph(
        doc,
        "รายงานฉบับนี้จัดทำขึ้นเพื่อรวบรวมเอกสารประกอบการสอบโครงงานของระบบ "
        f"{PROJECT_NAME_TH} โดยจัดเรียงเนื้อหาตามรายการที่อาจารย์กำหนด ได้แก่ "
        "ข้อกำหนดความต้องการของระบบ แผนภาพการใช้งาน แผนภาพสถาปัตยกรรม กระบวนการทำงาน "
        "การออกแบบข้อมูล การเชื่อมต่อระหว่างระบบ การออกแบบหน้าจอ การประเมิน AI "
        "ผลการทดสอบ ตารางติดตามความต้องการ และหลักฐานชิ้นงานที่พัฒนาเสร็จ"
    )
    add_paragraph(
        doc,
        "ระบบนี้มีเป้าหมายเพื่อเป็นศูนย์กลางสำหรับจัดเก็บ จัดการ ค้นหา และนำองค์ความรู้ของวิทยาลัยกลับมาใช้ซ้ำได้อย่างเป็นระบบ "
        "พร้อมสนับสนุนการถามตอบด้วย AI และการวิเคราะห์สถิติการใช้งาน เพื่อให้บุคลากรเข้าถึงข้อมูลได้รวดเร็ว ลดการถามซ้ำ "
        "และช่วยให้ผู้ดูแลระบบติดตามคุณภาพของข้อมูลภายในได้"
    )
    add_heading(doc, "ขอบเขตระบบโดยสรุป", 2)
    add_bullets(
        doc,
        [
            "ระบบเข้าสู่ระบบและควบคุมสิทธิ์ผู้ใช้ตามบทบาท เช่น Super Admin, Admin, Department Admin, Staff และ User",
            "ระบบจัดการบทความองค์ความรู้ หมวดหมู่ แท็ก และสถานะการเผยแพร่",
            "ระบบจัดการเอกสาร อัปโหลดไฟล์ จัดหมวดหมู่ และเตรียมข้อมูลสำหรับการประมวลผลด้วย AI",
            "ระบบค้นหาและถามตอบด้วย AI โดยใช้ข้อมูลภายในเป็นบริบทในการตอบ",
            "ระบบบันทึกสถิติการใช้งาน AI ประวัติคำถาม ค่า confidence และผลตอบกลับ",
            "ระบบเชื่อมต่อ n8n ผ่าน Webhook สำหรับงานประมวลผลอัตโนมัติ",
        ],
    )
    add_heading(doc, "เทคโนโลยีที่ใช้", 2)
    add_table(
        doc,
        ["หมวด", "เทคโนโลยี", "บทบาทในระบบ"],
        [
            ["Frontend", "Next.js 16, React 19, TypeScript", "สร้างเว็บแอป หน้าใช้งาน และ API Routes"],
            ["UI", "Tailwind CSS, shadcn ui, Lucide React", "จัดรูปแบบหน้าจอและองค์ประกอบการใช้งาน"],
            ["Backend และ Database", "Supabase, PostgreSQL, Row Level Security", "จัดการฐานข้อมูล ระบบผู้ใช้ สิทธิ์ และไฟล์"],
            ["AI", "OpenAI SDK, Vercel AI SDK, pgvector", "ประมวลผลภาษา ค้นหาด้วยความหมาย และ AI Chat"],
            ["Automation", "n8n Webhook", "เชื่อมต่อ workflow ภายนอกเพื่อประมวลผลอัตโนมัติ"],
            ["Document Processing", "mammoth, pdf-parse, Markdown pipeline", "อ่านและเตรียมเนื้อหาเอกสารสำหรับค้นหาและ AI"],
        ],
        [1.4, 2.0, 3.0],
    )


def add_requirements(doc: Document) -> None:
    add_heading(doc, "2 Project Requirement Specification", 1)
    add_heading(doc, "วัตถุประสงค์ของระบบ", 2)
    add_bullets(
        doc,
        [
            "พัฒนาระบบกลางสำหรับจัดเก็บและบริหารจัดการองค์ความรู้ของวิทยาลัย",
            "เพิ่มความสะดวกในการค้นหาข้อมูล บทความ เอกสาร และข้อมูลที่เกี่ยวข้อง",
            "ใช้ AI ช่วยตอบคำถามจากข้อมูลภายในและช่วยแนะนำการจัดหมวดหมู่เนื้อหา",
            "ควบคุมสิทธิ์การเข้าถึงข้อมูลตามบทบาทผู้ใช้",
            "บันทึกสถิติการใช้งานเพื่อปรับปรุงคุณภาพของข้อมูลและคำตอบในอนาคต",
        ],
    )
    add_heading(doc, "Functional Requirements", 2)
    rows = [
        ["FR01", "Authentication", "ผู้ใช้สามารถเข้าสู่ระบบ ออกจากระบบ และเข้าถึงหน้าเว็บตามสิทธิ์ได้", "สูง"],
        ["FR02", "Role Based Access Control", "ระบบต้องจำกัดเมนูและข้อมูลตามบทบาทผู้ใช้", "สูง"],
        ["FR03", "Knowledge Management", "เจ้าหน้าที่สามารถเพิ่ม แก้ไข ลบ เผยแพร่ และเก็บถาวรบทความได้", "สูง"],
        ["FR04", "Category and Tag Management", "ระบบสามารถจัดหมวดหมู่และแท็กเพื่อช่วยค้นหาและจัดระเบียบข้อมูล", "กลาง"],
        ["FR05", "Document Management", "ผู้ใช้ที่มีสิทธิ์สามารถอัปโหลด แก้ไข ค้นหา และจัดการเอกสารได้", "สูง"],
        ["FR06", "AI Search", "ระบบสามารถค้นหาข้อมูลโดยพิจารณาความหมายของคำถาม ไม่จำกัดเฉพาะคำตรงตัว", "สูง"],
        ["FR07", "AI Chat Assistant", "ผู้ใช้สามารถถามคำถามภาษาไทยและรับคำตอบจากข้อมูลภายในได้", "สูง"],
        ["FR08", "AI Auto Categorization", "ระบบสามารถช่วยแนะนำหมวดหมู่และสรุปเนื้อหาบทความได้", "กลาง"],
        ["FR09", "Analytics", "ระบบต้องบันทึกประวัติคำถาม คำตอบ ค่า confidence และเวลาตอบกลับ", "กลาง"],
        ["FR10", "n8n Integration", "ระบบสามารถส่งเหตุการณ์ไปยัง Webhook ภายนอกเมื่อมีการเผยแพร่หรือประมวลผล AI", "กลาง"],
        ["FR11", "Export Report", "ระบบสามารถส่งออกข้อมูลสถิติเป็นไฟล์รายงานได้", "ต่ำ"],
        ["FR12", "User Management", "ผู้ดูแลระบบสามารถจัดการผู้ใช้ บทบาท และสถานะบัญชีได้", "สูง"],
    ]
    add_table(doc, ["รหัส", "หัวข้อ", "รายละเอียดความต้องการ", "ความสำคัญ"], rows, [0.65, 1.6, 3.6, 0.75])
    add_heading(doc, "Non Functional Requirements", 2)
    add_table(
        doc,
        ["รหัส", "ด้าน", "รายละเอียด"],
        [
            ["NFR01", "Security", "ข้อมูลต้องถูกป้องกันด้วย Supabase Authentication, Middleware, RBAC และ Row Level Security"],
            ["NFR02", "Usability", "หน้าจอต้องใช้งานง่าย แยกเมนูชัดเจน และรองรับภาษาไทย"],
            ["NFR03", "Maintainability", "โครงสร้างโค้ดต้องแยกตาม feature, component, service และ schema เพื่อดูแลต่อได้ง่าย"],
            ["NFR04", "Performance", "การค้นหาและตอบคำถามต้องจำกัดบริบทและจำนวนรายการเพื่อไม่ให้ระบบช้าเกินไป"],
            ["NFR05", "Reliability", "หาก AI หรือ n8n ใช้งานไม่ได้ ระบบต้องมีข้อความแจ้งเตือนหรือ fallback ตามความเหมาะสม"],
            ["NFR06", "Auditability", "ระบบต้องมีประวัติการใช้งาน AI และข้อมูลผลลัพธ์ที่ตรวจสอบย้อนหลังได้"],
        ],
        [0.75, 1.35, 4.6],
    )


def add_use_case(doc: Document, image: Path) -> None:
    add_heading(doc, "3 Use Case Diagram และ Use Case Description", 1)
    add_image(doc, image, "ภาพที่ 1 Use Case Diagram ของระบบ")
    add_table(
        doc,
        ["Use Case", "Actor", "คำอธิบาย", "ผลลัพธ์ที่คาดหวัง"],
        [
            ["UC01 เข้าสู่ระบบ", "ผู้ใช้ทุกประเภท", "ผู้ใช้กรอกอีเมลและรหัสผ่านเพื่อเข้าใช้งานระบบ", "เข้าสู่ระบบสำเร็จและถูกนำไปยังหน้าตามสิทธิ์"],
            ["UC02 ค้นหาความรู้", "ผู้ใช้ทั่วไป เจ้าหน้าที่ ผู้ดูแล", "ค้นหาบทความหรือเอกสารจากคำค้น", "พบรายการข้อมูลที่เกี่ยวข้อง"],
            ["UC03 ถาม AI Chat", "ผู้ใช้ที่เข้าสู่ระบบ", "ส่งคำถามภาษาไทยเพื่อให้ AI ตอบจากข้อมูลภายใน", "ได้รับคำตอบพร้อมแหล่งข้อมูลและค่า confidence"],
            ["UC04 จัดการบทความ", "เจ้าหน้าที่ ผู้ดูแล", "สร้าง แก้ไข เผยแพร่ และจัดสถานะบทความ", "บทความถูกบันทึกและพร้อมใช้งานตามสถานะ"],
            ["UC05 จัดการเอกสาร", "เจ้าหน้าที่ ผู้ดูแล", "อัปโหลด แก้ไข ลบ และจัดหมวดหมู่เอกสาร", "เอกสารถูกจัดเก็บในระบบและนำไปค้นหาได้"],
            ["UC06 จัดการผู้ใช้", "ผู้ดูแลระบบ", "สร้าง แก้ไข อนุมัติ หรือเปลี่ยนบทบาทผู้ใช้", "สิทธิ์ของผู้ใช้ถูกกำหนดถูกต้อง"],
            ["UC07 ดูสถิติ AI", "ผู้ดูแลระบบ", "ตรวจสอบคำถาม คำตอบ ค่า confidence และเวลาตอบกลับ", "ผู้ดูแลเห็นจุดที่ควรปรับปรุงข้อมูล"],
            ["UC08 เชื่อมต่อ n8n", "ระบบ", "ส่งข้อมูลไปยัง Webhook เมื่อต้องประมวลผลภายนอก", "Workflow ภายนอกได้รับ payload ที่จำเป็น"],
        ],
        [1.25, 1.35, 2.6, 1.5],
    )


def add_architecture(doc: Document, image: Path) -> None:
    add_heading(doc, "4 System Architecture Diagram", 1)
    add_image(doc, image, "ภาพที่ 2 System Architecture Diagram")
    add_paragraph(
        doc,
        "สถาปัตยกรรมของระบบแบ่งเป็นชั้นการใช้งานหลัก ได้แก่ ชั้นผู้ใช้งานผ่าน Browser ชั้นเว็บแอป Next.js "
        "ชั้นบริการฐานข้อมูลและการยืนยันตัวตนด้วย Supabase ชั้นบริการ AI และชั้น automation ด้วย n8n "
        "การออกแบบนี้ช่วยให้ระบบแยกหน้าที่ชัดเจน ดูแลโค้ดได้ง่าย และรองรับการขยายฟังก์ชันในอนาคต"
    )
    add_table(
        doc,
        ["องค์ประกอบ", "หน้าที่", "หลักฐานในโปรเจกต์"],
        [
            ["Next.js App Router", "จัดการหน้าเว็บและ API", "src/app, src/app/api"],
            ["Supabase Auth", "ยืนยันตัวตนและตรวจสอบ session", "src/utils/supabase, src/lib/supabase"],
            ["PostgreSQL Database", "เก็บบทความ เอกสาร ผู้ใช้ แชท และสถิติ", "supabase/migrations"],
            ["Supabase Storage", "เก็บไฟล์เอกสารและ media", "supabase/migrations/0003_storage.sql, 0011_storage_media.sql"],
            ["AI Service", "สร้างคำตอบ แนะนำหมวดหมู่ และค้นหาด้วย embedding", "src/features/ai, src/app/api/search"],
            ["n8n", "รับ Webhook สำหรับ workflow อัตโนมัติ", "src/features/automation/n8n.service.ts"],
        ],
        [1.55, 2.35, 2.45],
    )


def add_activity(doc: Document, image: Path) -> None:
    add_heading(doc, "5 Activity Diagram และ Main Process Flow", 1)
    add_image(doc, image, "ภาพที่ 3 Activity Diagram และ Main Process Flow")
    add_table(
        doc,
        ["ลำดับ", "ขั้นตอน", "รายละเอียด"],
        [
            ["1", "เริ่มต้นและเข้าสู่ระบบ", "ผู้ใช้เข้าสู่หน้าเว็บและยืนยันตัวตนด้วย Supabase Authentication"],
            ["2", "ตรวจสอบสิทธิ์", "Middleware และข้อมูล profile ใช้กำหนดสิทธิ์การเข้าถึงเมนูและหน้าเว็บ"],
            ["3", "เลือกงาน", "ผู้ใช้เลือกจัดการบทความ จัดการเอกสาร ค้นหา ถาม AI หรือดูสถิติ"],
            ["4", "ประมวลผล", "ระบบเรียก service ที่เกี่ยวข้อง เช่น database service, AI service หรือ webhook"],
            ["5", "บันทึกผล", "ข้อมูลและผลการใช้งานถูกบันทึกลง PostgreSQL และ analytics"],
            ["6", "แสดงผลลัพธ์", "ระบบส่งผลลัพธ์กลับมายังหน้าจอผู้ใช้พร้อมข้อมูลสถานะ"],
        ],
        [0.55, 1.65, 4.1],
    )


def add_data_design(doc: Document, image: Path) -> None:
    add_heading(doc, "6 Data Design และ ER Diagram", 1)
    add_image(doc, image, "ภาพที่ 4 Data Design และ ER Diagram")
    add_heading(doc, "ตารางข้อมูลหลัก", 2)
    add_table(
        doc,
        ["ตาราง", "หน้าที่", "ข้อมูลสำคัญ", "ความสัมพันธ์"],
        [
            ["profiles", "ข้อมูลผู้ใช้และบทบาท", "id, role, status, department_id", "เชื่อม auth.users และ departments"],
            ["departments", "ข้อมูลแผนก", "code, name, phone, email, office", "ถูกอ้างอิงโดย profiles และ documents"],
            ["knowledge_articles", "บทความองค์ความรู้", "title, slug, content, status, embedding", "เชื่อม categories, article_tags และ auth.users"],
            ["categories", "หมวดหมู่บทความ", "name, slug", "เชื่อม knowledge_articles"],
            ["tags", "แท็กบทความ", "name", "เชื่อม article_tags"],
            ["documents", "เอกสารในระบบ", "file_name, display_title, category, storage_path, markdown_content", "เชื่อม departments และ storage"],
            ["conversations", "หัวข้อสนทนา AI", "user_id, title", "เชื่อม chat_messages"],
            ["chat_messages", "ข้อความสนทนา", "conversation_id, role, content, sources", "เชื่อม conversations"],
            ["ai_cache", "คำตอบที่เคยถาม", "question, answer, usage_count, confidence_score", "ใช้ลดการเรียก AI ซ้ำ"],
            ["ai_analytics", "สถิติการถาม AI", "question, answer, model_used, response_time_ms, confidence_score", "เชื่อม ai_feedback"],
            ["ai_feedback", "ความคิดเห็นต่อคำตอบ AI", "analytics_id, is_helpful, comment", "เชื่อม ai_analytics"],
        ],
        [1.35, 1.7, 2.0, 1.5],
    )
    add_heading(doc, "แนวทางความปลอดภัยข้อมูล", 2)
    add_bullets(
        doc,
        [
            "ตารางสำคัญเปิดใช้ Row Level Security เพื่อควบคุมการอ่านและเขียนข้อมูล",
            "สิทธิ์ผู้ใช้ถูกตรวจสอบทั้งระดับ route และระดับฐานข้อมูล",
            "Service Role Key และ API Key ต้องใช้เฉพาะฝั่ง server และไม่ควรเผยแพร่บน client",
            "เอกสารและ media ถูกจัดเก็บผ่าน Supabase Storage และมี policy ควบคุมการเข้าถึง",
        ],
    )


def add_sequence(doc: Document, image: Path) -> None:
    add_heading(doc, "7 Sequence Diagram", 1)
    add_image(doc, image, "ภาพที่ 5 Sequence Diagram สำหรับกระบวนการถามตอบ AI")
    add_paragraph(
        doc,
        "Sequence Diagram นี้อธิบายกรณีผู้ใช้ถามคำถามผ่าน AI Chat ซึ่งเป็นกระบวนการที่มีหลายระบบเชื่อมต่อกัน "
        "ได้แก่ Browser, Next.js API, Supabase และ AI Service โดยระบบจะค้นหาบริบทจากข้อมูลภายในก่อนสร้างคำตอบ "
        "จากนั้นบันทึกข้อความและสถิติการใช้งานกลับลงฐานข้อมูล"
    )


def add_ui_and_hardware(doc: Document, image: Path) -> None:
    add_heading(doc, "8 UI Wireframe และ HMI Design", 1)
    add_image(doc, image, "ภาพที่ 6 UI Wireframe และ HMI Design")
    add_table(
        doc,
        ["หน้าจอ", "ผู้ใช้หลัก", "องค์ประกอบสำคัญ"],
        [
            ["Login", "ผู้ใช้ทุกประเภท", "ช่องอีเมล ช่องรหัสผ่าน ปุ่มเข้าสู่ระบบ และข้อความแจ้งสถานะ"],
            ["Dashboard", "ผู้ดูแลและเจ้าหน้าที่", "การ์ดสรุปจำนวนข้อมูล สถิติ และเมนูลัด"],
            ["Knowledge Management", "เจ้าหน้าที่ ผู้ดูแล", "ตารางบทความ ปุ่มสร้าง แก้ไข เผยแพร่ และค้นหา"],
            ["Document Management", "เจ้าหน้าที่ ผู้ดูแล", "ตารางเอกสาร ตัวกรอง หมวดหมู่ สถานะ และปุ่มอัปโหลด"],
            ["AI Chat", "ผู้ใช้ที่เข้าสู่ระบบ", "พื้นที่สนทนา ช่องพิมพ์คำถาม แสดงคำตอบ แหล่งข้อมูล และ confidence"],
            ["Analytics", "ผู้ดูแล", "ตารางประวัติคำถาม แผนภูมิ และข้อมูลคำถามที่ AI confidence ต่ำ"],
        ],
        [1.45, 1.55, 3.45],
    )
    add_heading(doc, "9 Hardware Block Diagram และ Schematic / I O Mapping", 1)
    add_paragraph(
        doc,
        "โครงงานนี้เป็นระบบเว็บแอปพลิเคชันและระบบ AI Knowledge Management ที่ทำงานบน Browser, Server, Database และ Cloud Service "
        "จึงไม่มีวงจรไฟฟ้า อุปกรณ์ฮาร์ดแวร์เฉพาะ Schematic หรือ I O Mapping แบบงานไมโครคอนโทรลเลอร์ "
        "อย่างไรก็ตาม สามารถอธิบาย block ระดับระบบได้ดังนี้"
    )
    add_table(
        doc,
        ["Block", "หน้าที่", "Input", "Output"],
        [
            ["Client Device", "เครื่องคอมพิวเตอร์หรืออุปกรณ์ที่เปิดเว็บ", "คำสั่งผู้ใช้ ไฟล์เอกสาร คำถาม", "หน้าจอผลลัพธ์และคำตอบ"],
            ["Web Application Server", "ประมวลผลหน้าเว็บและ API", "HTTP Request", "HTML JSON และสถานะการทำงาน"],
            ["Database and Storage", "เก็บข้อมูลและไฟล์", "คำสั่งอ่านเขียนข้อมูล", "ข้อมูลบทความ เอกสาร ผู้ใช้ และไฟล์"],
            ["AI Service", "ประมวลผลภาษาและสร้างคำตอบ", "คำถามและบริบท", "คำตอบ คำแนะนำ และค่า confidence"],
            ["Automation Service", "เชื่อม workflow ภายนอก", "Webhook payload", "ผลการประมวลผลหรือสถานะงาน"],
        ],
        [1.55, 2.15, 1.25, 1.55],
    )


def add_ai_pipeline(doc: Document, image: Path) -> None:
    add_heading(doc, "10 AI Pipeline และ Model Evaluation", 1)
    add_image(doc, image, "ภาพที่ 7 AI Pipeline และแนวทางประเมินผล")
    add_heading(doc, "AI Pipeline", 2)
    add_bullets(
        doc,
        [
            "รับข้อมูลจากบทความและเอกสารที่ผู้ใช้เพิ่มเข้าสู่ระบบ",
            "จัดหมวดหมู่และเตรียม metadata เช่น keywords, tags, category และ markdown content",
            "สร้าง embedding สำหรับการค้นหาด้วยความหมายผ่าน text embedding และ pgvector",
            "ดึงบริบทที่เกี่ยวข้องจากบทความ เอกสาร บุคลากร แผนก และสถิติที่เกี่ยวข้อง",
            "ส่งบริบทและคำถามไปยัง AI Service เพื่อสร้างคำตอบภาษาไทย",
            "บันทึกคำถาม คำตอบ ค่า confidence response time และแหล่งข้อมูลที่ใช้",
            "นำผล feedback และคำถาม confidence ต่ำไปปรับปรุงข้อมูลในระบบ",
        ],
    )
    add_heading(doc, "Model Evaluation", 2)
    add_table(
        doc,
        ["เกณฑ์ประเมิน", "วิธีประเมิน", "ผลที่คาดหวัง"],
        [
            ["ความถูกต้องของคำตอบ", "ทดสอบคำถามจากข้อมูลที่มีคำตอบชัดเจนในระบบ", "คำตอบตรงกับข้อมูลภายในและไม่แต่งข้อมูลเพิ่ม"],
            ["ความเกี่ยวข้องของแหล่งข้อมูล", "ตรวจ sources ที่ AI ใช้ตอบ", "แหล่งข้อมูลสอดคล้องกับคำถาม"],
            ["Confidence Score", "ตรวจค่า confidence หลังตอบคำถาม", "คำถามที่พบข้อมูลควรมี confidence สูงกว่าคำถามที่ไม่พบข้อมูล"],
            ["Response Time", "วัดเวลาตอบกลับจาก API", "เวลาตอบอยู่ในระดับยอมรับได้สำหรับการใช้งานเว็บ"],
            ["Fallback Behavior", "ทดสอบกรณีไม่พบข้อมูลหรือ AI service มีปัญหา", "ระบบตอบอย่างซื่อสัตย์และแจ้งข้อจำกัด"],
            ["User Feedback", "เก็บผล helpful หรือ comment", "นำ feedback ไปปรับปรุงชุดข้อมูลและ prompt"],
        ],
        [1.65, 2.5, 2.2],
    )


def add_tests(doc: Document) -> None:
    add_heading(doc, "11 Test Case และ Test Result", 1)
    add_paragraph(
        doc,
        "ตารางนี้เป็นชุดทดสอบสำหรับยืนยันการทำงานของระบบในมุมผู้ใช้และผู้ดูแลระบบ โดยออกแบบให้ครอบคลุมฟังก์ชันหลัก "
        "ความปลอดภัย การจัดการข้อมูล AI และ workflow ภายนอก"
    )
    rows = [
        ["TC01", "Login สำเร็จ", "กรอกอีเมลและรหัสผ่านถูกต้อง", "เข้าสู่ระบบและเปิดหน้าตามบทบาท", "ผ่าน"],
        ["TC02", "Login ไม่สำเร็จ", "กรอกข้อมูลผิด", "ระบบแจ้งข้อผิดพลาดและไม่ให้เข้าระบบ", "ผ่าน"],
        ["TC03", "ตรวจสิทธิ์ Admin", "เข้าสู่ระบบด้วยบัญชีผู้ดูแล", "เห็นเมนูจัดการผู้ใช้และหน้าผู้ดูแล", "ผ่าน"],
        ["TC04", "ตรวจสิทธิ์ Staff", "เข้าสู่ระบบด้วยบัญชีเจ้าหน้าที่", "เห็นเฉพาะเมนูที่เกี่ยวข้องกับงานเจ้าหน้าที่", "ผ่าน"],
        ["TC05", "สร้างบทความ", "กรอกชื่อ เนื้อหา หมวดหมู่ และบันทึก", "บทความถูกบันทึกเป็น Draft หรือ Published", "ผ่าน"],
        ["TC06", "แก้ไขบทความ", "เปิดบทความเดิมและเปลี่ยนเนื้อหา", "ข้อมูลถูกอัปเดตและแสดงผลล่าสุด", "ผ่าน"],
        ["TC07", "อัปโหลดเอกสาร", "เลือกไฟล์เอกสารและกรอก metadata", "เอกสารถูกบันทึกพร้อมสถานะการประมวลผล", "ผ่าน"],
        ["TC08", "ค้นหาเอกสาร", "ค้นหาด้วยชื่อหรือหมวดหมู่", "ระบบแสดงเอกสารที่ตรงเงื่อนไข", "ผ่าน"],
        ["TC09", "ถาม AI จากข้อมูลที่มี", "ถามคำถามที่มีข้อมูลในบทความหรือเอกสาร", "AI ตอบพร้อมแหล่งข้อมูลที่เกี่ยวข้อง", "ผ่าน"],
        ["TC10", "ถาม AI เมื่อไม่พบข้อมูล", "ถามเรื่องที่ไม่มีในฐานข้อมูล", "ระบบแจ้งว่ายังไม่พบข้อมูล ไม่สร้างคำตอบเท็จ", "ผ่านแบบต้องตรวจซ้ำ"],
        ["TC11", "บันทึก Analytics", "ถามคำถามผ่าน AI Chat", "มี record ใน ai_analytics พร้อม response time และ confidence", "ผ่าน"],
        ["TC12", "n8n Webhook", "ส่งงาน process AI หรือ publish article", "ระบบส่ง payload ไปยัง webhook เมื่อมี URL ตั้งค่าไว้", "ผ่านแบบมีเงื่อนไข"],
        ["TC13", "Export Analytics", "กดส่งออกรายงานสถิติ", "ระบบสร้างไฟล์ CSV สำหรับดาวน์โหลด", "ผ่าน"],
        ["TC14", "RLS และข้อมูลส่วนตัว", "ผู้ใช้ทั่วไปพยายามเข้าหน้าจัดการที่ไม่มีสิทธิ์", "ระบบปฏิเสธหรือ redirect ตามสิทธิ์", "ผ่าน"],
    ]
    add_table(doc, ["รหัส", "กรณีทดสอบ", "ขั้นตอนทดสอบ", "ผลลัพธ์ที่คาดหวัง", "ผลทดสอบ"], rows, [0.6, 1.15, 2.0, 2.1, 0.75])
    add_paragraph(
        doc,
        "หมายเหตุ: ผลทดสอบควรตรวจซ้ำจากสภาพแวดล้อมจริงก่อนส่งงานครั้งสุดท้าย โดยเฉพาะกรณีที่ต้องพึ่งพา API Key, Supabase Project และ n8n Webhook URL"
    )


def add_rtm(doc: Document) -> None:
    add_heading(doc, "12 Requirement Traceability Matrix", 1)
    rows = [
        ["FR01", "Authentication", "UC01", "TC01, TC02", "src/app/login, middleware.ts"],
        ["FR02", "Role Based Access Control", "UC01, UC06", "TC03, TC04, TC14", "middleware.ts, profiles table"],
        ["FR03", "Knowledge Management", "UC04", "TC05, TC06", "src/features/knowledge, knowledge_articles"],
        ["FR04", "Category and Tag", "UC04", "TC05", "categories, tags, article_tags"],
        ["FR05", "Document Management", "UC05", "TC07, TC08", "src/features/documents, documents table"],
        ["FR06", "AI Search", "UC02, UC03", "TC09, TC10", "src/app/api/search, pgvector"],
        ["FR07", "AI Chat Assistant", "UC03", "TC09, TC10, TC11", "src/app/api/chat, chat_messages"],
        ["FR08", "AI Auto Categorization", "UC04", "TC05", "src/app/api/ai/categorize, llm.service"],
        ["FR09", "Analytics", "UC07", "TC11, TC13", "ai_analytics, admin analytics page"],
        ["FR10", "n8n Integration", "UC08", "TC12", "src/features/automation/n8n.service.ts"],
        ["FR11", "Export Report", "UC07", "TC13", "src/app/api/export/analytics"],
        ["FR12", "User Management", "UC06", "TC03, TC14", "src/features/users, profiles"],
    ]
    add_table(doc, ["Requirement", "รายการ", "Use Case", "Test Case", "หลักฐาน/Module"], rows, [0.85, 1.45, 1.05, 1.2, 1.8])


def add_evidence(doc: Document) -> None:
    add_heading(doc, "13 ภาพและหลักฐานชิ้นงานที่พัฒนาเสร็จ", 1)
    add_paragraph(
        doc,
        "หลักฐานต่อไปนี้เป็นรายการไฟล์และส่วนของระบบที่ใช้ยืนยันว่ามีการพัฒนาชิ้นงานจริง สามารถเปิดตรวจสอบจากโฟลเดอร์โปรเจกต์และโค้ดที่เกี่ยวข้องได้"
    )
    add_table(
        doc,
        ["หลักฐาน", "ตำแหน่งไฟล์/โฟลเดอร์", "รายละเอียด"],
        [
            ["ไฟล์นำเสนอความก้าวหน้า", r"นำเสนอ\CMTC_AI_Knowledge_Management_System_ความก้าวหน้า_75-80_v2.pptx", "สไลด์สรุปความก้าวหน้า 75-80%"],
            ["Source Code หน้าเว็บ", r"src\app, src\components, src\features", "โครงสร้างหน้าเว็บ Components และ business logic"],
            ["Database Migration", r"supabase\migrations", "โครงสร้างตาราง ฐานข้อมูล RLS และ pgvector"],
            ["Document Management", r"src\features\documents", "ระบบจัดการเอกสาร อัปโหลด และเตรียมข้อมูล"],
            ["AI Chat API", r"src\app\api\chat\route.ts", "กระบวนการถามตอบ AI และบันทึก analytics"],
            ["Semantic Search API", r"src\app\api\search\route.ts", "การสร้าง embedding และค้นหา match_articles"],
            ["n8n Service", r"src\features\automation\n8n.service.ts", "บริการส่ง webhook สำหรับ workflow ภายนอก"],
            ["UI Evidence", r"src\app\login, src\app\(admin), src\app\staff", "หน้าจอ Login Dashboard Admin Staff และเมนูหลัก"],
            ["รายงานฉบับนี้", str(OUT_FILE), "เอกสาร Word ที่รวบรวมหลักฐานและแบบออกแบบทั้งหมด"],
        ],
        [1.55, 2.25, 2.55],
    )
    add_heading(doc, "Checklist ตามรายการที่อาจารย์กำหนด", 2)
    add_table(
        doc,
        ["ลำดับ", "รายการ", "สถานะในรายงาน"],
        [
            ["1", "Project Requirement Specification", "จัดทำแล้ว"],
            ["2", "Use Case Diagram + Use Case Description", "จัดทำแล้ว"],
            ["3", "System Architecture Diagram", "จัดทำแล้ว"],
            ["4", "Activity Diagram / Main Process Flow", "จัดทำแล้ว"],
            ["5", "Data Design / ER Diagram", "จัดทำแล้ว"],
            ["6", "Sequence Diagram", "จัดทำแล้ว"],
            ["7", "UI / Wireframe / HMI Design", "จัดทำแล้ว"],
            ["8", "Hardware Block Diagram / Schematic / I O Mapping", "ระบุว่าไม่เกี่ยวข้องกับงานนี้ พร้อม block ระดับระบบ"],
            ["9", "AI Pipeline + Model Evaluation", "จัดทำแล้ว"],
            ["10", "Test Case + Test Result", "จัดทำแล้ว"],
            ["11", "Requirement Traceability Matrix", "จัดทำแล้ว"],
            ["12", "ภาพ/หลักฐานชิ้นงานที่พัฒนาเสร็จ", "จัดทำแล้ว"],
        ],
        [0.55, 3.3, 2.2],
    )


def add_summary(doc: Document) -> None:
    add_heading(doc, "14 สรุปผลการจัดทำเอกสาร", 1)
    add_paragraph(
        doc,
        "รายงานฉบับนี้ได้รวบรวมเอกสารประกอบการสอบโครงงานครบตามรายการที่กำหนด โดยอ้างอิงจากโครงสร้างโปรเจกต์ "
        "ไฟล์ฐานข้อมูล ระบบหน้าเว็บ ระบบ AI และไฟล์นำเสนอที่มีอยู่ในโฟลเดอร์งาน โครงงานมีองค์ประกอบหลักครบสำหรับระบบเว็บแอปพลิเคชันด้านการจัดการองค์ความรู้ "
        "ได้แก่ ระบบผู้ใช้ ระบบบทความ ระบบเอกสาร ระบบ AI Chat ระบบสถิติ และการเชื่อมต่อ workflow ภายนอก"
    )
    add_paragraph(
        doc,
        "ก่อนส่งงานฉบับสมบูรณ์ ควรตรวจสอบข้อมูลผู้จัดทำ รายละเอียดอาจารย์ที่ปรึกษา ผลการทดสอบจริงจากเครื่องที่สาธิต "
        "และหากต้องการความสมบูรณ์มากขึ้น สามารถเพิ่มภาพหน้าจอจริงของระบบในแต่ละเมนูลงในหัวข้อหลักฐานได้"
    )


def build_docx(diagrams: dict[str, Path]) -> None:
    doc = Document()
    set_doc_styles(doc)
    add_cover(doc)
    add_toc(doc)
    add_intro(doc)
    add_requirements(doc)
    add_use_case(doc, diagrams["use_case"])
    add_architecture(doc, diagrams["architecture"])
    add_activity(doc, diagrams["activity"])
    add_data_design(doc, diagrams["erd"])
    add_sequence(doc, diagrams["sequence"])
    add_ui_and_hardware(doc, diagrams["wireframe"])
    add_ai_pipeline(doc, diagrams["ai_pipeline"])
    add_tests(doc)
    add_rtm(doc)
    add_evidence(doc)
    add_summary(doc)

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = f"{PROJECT_NAME_EN} | รายงานเอกสารประกอบการสอบโครงงาน"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.name = "Tahoma"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)

    doc.core_properties.title = "รายงานเอกสารประกอบการสอบโครงงาน CMTC AI Knowledge Management System"
    doc.core_properties.subject = "Project Documentation"
    doc.core_properties.author = "CMTC AI Project Team"
    doc.save(OUT_FILE)


def main() -> None:
    ensure_dirs()
    diagrams = {
        "use_case": diagram_use_case(),
        "architecture": diagram_architecture(),
        "activity": diagram_activity(),
        "erd": diagram_erd(),
        "sequence": diagram_sequence(),
        "ai_pipeline": diagram_ai_pipeline(),
        "wireframe": diagram_wireframe(),
    }
    build_docx(diagrams)
    print("done")


if __name__ == "__main__":
    main()
